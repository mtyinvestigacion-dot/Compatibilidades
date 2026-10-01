from datetime import datetime
from io import BytesIO

from flask import Blueprint, abort, flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required, login_user, logout_user
from openpyxl import Workbook, load_workbook

from app import db
from app.auth import admin_required
from app.models import RESULTADOS, ROLES, Compatibilidad, Usuario

main = Blueprint('main', __name__)

COLUMNAS_IMPORTACION = (
    'panel_reemplazo', 'marca_tv', 'modelo_tv', 'panel_original_tv',
    'pulgada', 'mainboard', 't_con', 'observaciones', 'resultado',
)
ENCABEZADOS_IMPORTACION = (
    'Panel de reemplazo', 'Marca TV', 'Modelo TV', 'Panel original TV',
    'Pulgada', 'Mainboard', 'T-Con', 'Observaciones', 'Resultado (Exitosa/Fallida)',
)
RUTA_LISTADO = {'exitosa': 'main.listar_exitosas', 'fallida': 'main.listar_fallidas'}

RUTAS_SIN_LOGIN = ('main.login', 'static', 'main.static')


@main.before_request
def require_login():
    if request.endpoint in RUTAS_SIN_LOGIN:
        return
    if not current_user.is_authenticated:
        return redirect(url_for('main.login', next=request.path))


def _validar_compatibilidad(form):
    errores = []
    if not form.get('panel_reemplazo', '').strip():
        errores.append('El panel de reemplazo es obligatorio.')
    if not form.get('marca_tv', '').strip():
        errores.append('La marca de la TV es obligatoria.')
    if not form.get('modelo_tv', '').strip():
        errores.append('El modelo de la TV es obligatorio.')
    if not form.get('panel_original_tv', '').strip():
        errores.append('El panel original de la TV es obligatorio.')
    if form.get('resultado') not in RESULTADOS:
        errores.append('El resultado debe ser Exitosa o Fallida.')
    return errores


# --- Auth -------------------------------------------------------------

@main.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('main.index'))

    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        usuario = Usuario.query.filter_by(username=username, activo=True).first()
        if usuario and usuario.check_password(password):
            login_user(usuario)
            next_url = request.args.get('next')
            return redirect(next_url or url_for('main.index'))
        flash('Usuario o contraseña incorrectos.', 'error')

    return render_template('login.html')


@main.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('main.login'))


# --- Compatibilidades ---------------------------------------------------

@main.route('/')
def index():
    return redirect(url_for('main.listar_exitosas'))


def _listar_compatibilidades(resultado):
    q = request.args.get('q', '').strip()
    marca = request.args.get('marca', '').strip()

    query = Compatibilidad.query.filter_by(activo=True, resultado=resultado)
    if marca:
        query = query.filter(Compatibilidad.marca_tv == marca)
    if q:
        like = f'%{q}%'
        query = query.filter(
            db.or_(
                Compatibilidad.panel_reemplazo.ilike(like),
                Compatibilidad.marca_tv.ilike(like),
                Compatibilidad.modelo_tv.ilike(like),
                Compatibilidad.panel_original_tv.ilike(like),
            )
        )
    compatibilidades = query.order_by(Compatibilidad.marca_tv, Compatibilidad.modelo_tv).all()

    marcas = [
        m[0] for m in
        db.session.query(Compatibilidad.marca_tv)
        .filter_by(activo=True, resultado=resultado)
        .distinct().order_by(Compatibilidad.marca_tv)
    ]

    return render_template(
        'compatibilidades/listar.html',
        compatibilidades=compatibilidades,
        q=q,
        marca=marca,
        marcas=marcas,
        resultado=resultado,
    )


@main.route('/compatibilidades/exitosas')
def listar_exitosas():
    return _listar_compatibilidades('exitosa')


@main.route('/compatibilidades/fallidas')
def listar_fallidas():
    return _listar_compatibilidades('fallida')


@main.route('/compatibilidades/nueva', methods=['GET', 'POST'])
@admin_required
def nueva_compatibilidad():
    resultado_default = request.args.get('resultado', 'exitosa')
    if resultado_default not in RESULTADOS:
        resultado_default = 'exitosa'

    if request.method == 'POST':
        errores = _validar_compatibilidad(request.form)
        if errores:
            for error in errores:
                flash(error, 'error')
            return render_template('compatibilidades/form.html', compatibilidad=None, valores=request.form)

        resultado = request.form.get('resultado')
        compatibilidad = Compatibilidad(
            panel_reemplazo=request.form.get('panel_reemplazo').strip(),
            marca_tv=request.form.get('marca_tv').strip().upper(),
            modelo_tv=request.form.get('modelo_tv').strip(),
            panel_original_tv=request.form.get('panel_original_tv').strip(),
            pulgada=request.form.get('pulgada', '').strip() or None,
            mainboard=request.form.get('mainboard', '').strip() or None,
            t_con=request.form.get('t_con', '').strip() or None,
            observaciones=request.form.get('observaciones', '').strip() or None,
            resultado=resultado,
            creado_por_id=current_user.id,
        )
        db.session.add(compatibilidad)
        db.session.commit()
        flash('Compatibilidad guardada correctamente.', 'success')
        return redirect(url_for(RUTA_LISTADO[resultado]))

    return render_template('compatibilidades/form.html', compatibilidad=None, valores={'resultado': resultado_default})


@main.route('/compatibilidades/<int:compat_id>/editar', methods=['GET', 'POST'])
@admin_required
def editar_compatibilidad(compat_id):
    compatibilidad = Compatibilidad.query.get_or_404(compat_id)

    if request.method == 'POST':
        errores = _validar_compatibilidad(request.form)
        if errores:
            for error in errores:
                flash(error, 'error')
            return render_template('compatibilidades/form.html', compatibilidad=compatibilidad, valores=request.form)

        compatibilidad.panel_reemplazo = request.form.get('panel_reemplazo').strip()
        compatibilidad.marca_tv = request.form.get('marca_tv').strip().upper()
        compatibilidad.modelo_tv = request.form.get('modelo_tv').strip()
        compatibilidad.panel_original_tv = request.form.get('panel_original_tv').strip()
        compatibilidad.pulgada = request.form.get('pulgada', '').strip() or None
        compatibilidad.mainboard = request.form.get('mainboard', '').strip() or None
        compatibilidad.t_con = request.form.get('t_con', '').strip() or None
        compatibilidad.observaciones = request.form.get('observaciones', '').strip() or None
        compatibilidad.resultado = request.form.get('resultado')
        compatibilidad.modificado_por_id = current_user.id
        compatibilidad.fecha_modificacion = datetime.utcnow()
        db.session.commit()
        flash('Compatibilidad actualizada correctamente.', 'success')
        return redirect(url_for(RUTA_LISTADO[compatibilidad.resultado]))

    valores = {
        'panel_reemplazo': compatibilidad.panel_reemplazo,
        'marca_tv': compatibilidad.marca_tv,
        'modelo_tv': compatibilidad.modelo_tv,
        'panel_original_tv': compatibilidad.panel_original_tv,
        'pulgada': compatibilidad.pulgada or '',
        'mainboard': compatibilidad.mainboard or '',
        't_con': compatibilidad.t_con or '',
        'observaciones': compatibilidad.observaciones or '',
        'resultado': compatibilidad.resultado,
    }
    return render_template('compatibilidades/form.html', compatibilidad=compatibilidad, valores=valores)


@main.route('/compatibilidades/<int:compat_id>/eliminar', methods=['POST'])
@admin_required
def eliminar_compatibilidad(compat_id):
    compatibilidad = Compatibilidad.query.get_or_404(compat_id)
    compatibilidad.activo = False
    db.session.commit()
    flash('Compatibilidad eliminada.', 'success')
    return redirect(url_for(RUTA_LISTADO[compatibilidad.resultado]))


@main.route('/compatibilidades/importar/plantilla')
@admin_required
def plantilla_importacion():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Compatibilidades'
    ws.append(ENCABEZADOS_IMPORTACION)
    ws.append(['PAN-12345', 'SAMSUNG', 'UN50TU7000', 'PANEL-ORIG-001', '50', 'MAIN-001', 'TCON-001', 'Opcional', 'Exitosa'])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name='plantilla_compatibilidades.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@main.route('/compatibilidades/importar', methods=['GET', 'POST'])
@admin_required
def importar_compatibilidades():
    if request.method == 'POST':
        archivo = request.files.get('archivo')
        if not archivo or not archivo.filename:
            flash('Selecciona un archivo Excel (.xlsx) para importar.', 'error')
            return render_template('compatibilidades/importar.html', resumen=None)

        if not archivo.filename.lower().endswith('.xlsx'):
            flash('El archivo debe tener formato .xlsx.', 'error')
            return render_template('compatibilidades/importar.html', resumen=None)

        try:
            wb = load_workbook(archivo, read_only=True, data_only=True)
            ws = wb.active
        except Exception:
            flash('No se pudo leer el archivo. Verifica que sea un Excel válido.', 'error')
            return render_template('compatibilidades/importar.html', resumen=None)

        filas = list(ws.iter_rows(values_only=True))
        wb.close()

        if not filas:
            flash('El archivo está vacío.', 'error')
            return render_template('compatibilidades/importar.html', resumen=None)

        creadas = 0
        duplicadas = []
        errores = []
        existentes = {
            (c.panel_reemplazo.strip().upper(), c.marca_tv.strip().upper(),
             c.modelo_tv.strip().upper(), c.panel_original_tv.strip().upper())
            for c in Compatibilidad.query.filter_by(activo=True).all()
        }

        for num_fila, fila in enumerate(filas[1:], start=2):
            if fila is None or all(c is None or str(c).strip() == '' for c in fila):
                continue

            valores = dict(zip(COLUMNAS_IMPORTACION, (fila or ()) + (None,) * len(COLUMNAS_IMPORTACION)))
            panel_reemplazo = str(valores.get('panel_reemplazo') or '').strip()
            marca_tv = str(valores.get('marca_tv') or '').strip().upper()
            modelo_tv = str(valores.get('modelo_tv') or '').strip()
            panel_original_tv = str(valores.get('panel_original_tv') or '').strip()

            if not panel_reemplazo or not marca_tv or not modelo_tv or not panel_original_tv:
                errores.append(f'Fila {num_fila}: faltan datos obligatorios (panel, marca, modelo o panel original).')
                continue

            clave = (panel_reemplazo.upper(), marca_tv.upper(), modelo_tv.upper(), panel_original_tv.upper())
            if clave in existentes:
                duplicadas.append(f'Fila {num_fila}: {panel_reemplazo} / {marca_tv} {modelo_tv} ya existe.')
                continue

            pulgada = str(valores.get('pulgada')).strip() if valores.get('pulgada') not in (None, '') else None
            mainboard = str(valores.get('mainboard')).strip() if valores.get('mainboard') not in (None, '') else None
            t_con = str(valores.get('t_con')).strip() if valores.get('t_con') not in (None, '') else None
            observaciones = str(valores.get('observaciones')).strip() if valores.get('observaciones') not in (None, '') else None

            resultado_raw = str(valores.get('resultado') or '').strip().lower()
            if not resultado_raw:
                resultado = 'exitosa'
            elif resultado_raw in ('exitosa', 'exito', 'éxito', 'ok'):
                resultado = 'exitosa'
            elif resultado_raw in ('fallida', 'fallo', 'falla', 'no'):
                resultado = 'fallida'
            else:
                errores.append(f'Fila {num_fila}: resultado "{valores.get("resultado")}" inválido (usa Exitosa o Fallida).')
                continue

            compatibilidad = Compatibilidad(
                panel_reemplazo=panel_reemplazo,
                marca_tv=marca_tv,
                modelo_tv=modelo_tv,
                panel_original_tv=panel_original_tv,
                pulgada=pulgada,
                mainboard=mainboard,
                t_con=t_con,
                observaciones=observaciones,
                resultado=resultado,
                creado_por_id=current_user.id,
            )
            db.session.add(compatibilidad)
            existentes.add(clave)
            creadas += 1

        if creadas:
            db.session.commit()
        else:
            db.session.rollback()

        resumen = {'creadas': creadas, 'duplicadas': duplicadas, 'errores': errores}
        if creadas:
            flash(f'{creadas} compatibilidad(es) importada(s) correctamente.', 'success')
        return render_template('compatibilidades/importar.html', resumen=resumen)

    return render_template('compatibilidades/importar.html', resumen=None)


@main.route('/api/marcas')
def api_marcas():
    marcas = [m[0] for m in db.session.query(Compatibilidad.marca_tv).filter_by(activo=True).distinct()]
    return jsonify(sorted(marcas))


@main.route('/api/modelos')
def api_modelos():
    modelos = [m[0] for m in db.session.query(Compatibilidad.modelo_tv).filter_by(activo=True).distinct()]
    return jsonify(sorted(modelos))


# --- Usuarios (solo admin) ----------------------------------------------

@main.route('/usuarios')
@admin_required
def listar_usuarios():
    usuarios = Usuario.query.order_by(Usuario.username).all()
    return render_template('usuarios/listar.html', usuarios=usuarios)


@main.route('/usuarios/nuevo', methods=['GET', 'POST'])
@admin_required
def nuevo_usuario():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        password_confirm = request.form.get('password_confirm', '')
        rol = request.form.get('rol', 'viewer')

        errores = []
        if not username:
            errores.append('El usuario es obligatorio.')
        elif Usuario.query.filter_by(username=username).first():
            errores.append('Ya existe un usuario con ese nombre.')
        if len(password) < 6:
            errores.append('La contraseña debe tener al menos 6 caracteres.')
        if password != password_confirm:
            errores.append('Las contraseñas no coinciden.')
        if rol not in ROLES:
            errores.append('Rol inválido.')

        if errores:
            for error in errores:
                flash(error, 'error')
            return render_template('usuarios/form.html', form=request.form)

        usuario = Usuario(username=username, rol=rol)
        usuario.set_password(password)
        db.session.add(usuario)
        db.session.commit()
        flash(f'Usuario "{username}" creado correctamente.', 'success')
        return redirect(url_for('main.listar_usuarios'))

    return render_template('usuarios/form.html', form=None)


@main.route('/usuarios/<int:usuario_id>/toggle', methods=['POST'])
@admin_required
def toggle_usuario(usuario_id):
    usuario = Usuario.query.get_or_404(usuario_id)
    if usuario.id == current_user.id:
        flash('No puedes desactivar tu propia cuenta.', 'error')
        return redirect(url_for('main.listar_usuarios'))

    usuario.activo = not usuario.activo
    db.session.commit()
    estado = 'activado' if usuario.activo else 'desactivado'
    flash(f'Usuario "{usuario.username}" {estado}.', 'success')
    return redirect(url_for('main.listar_usuarios'))


@main.route('/usuarios/<int:usuario_id>/reset-password', methods=['GET', 'POST'])
@admin_required
def reset_password(usuario_id):
    usuario = Usuario.query.get_or_404(usuario_id)

    if request.method == 'POST':
        password = request.form.get('password', '')
        password_confirm = request.form.get('password_confirm', '')

        errores = []
        if len(password) < 6:
            errores.append('La contraseña debe tener al menos 6 caracteres.')
        if password != password_confirm:
            errores.append('Las contraseñas no coinciden.')

        if errores:
            for error in errores:
                flash(error, 'error')
            return render_template('usuarios/reset_password.html', usuario=usuario)

        usuario.set_password(password)
        db.session.commit()
        flash(f'Contraseña de "{usuario.username}" actualizada.', 'success')
        return redirect(url_for('main.listar_usuarios'))

    return render_template('usuarios/reset_password.html', usuario=usuario)
